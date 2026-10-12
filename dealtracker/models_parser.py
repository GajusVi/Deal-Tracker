"""Read uploaded financial models (xlsx/xlsm) and surface headline outputs."""
import os
import re

KEY_OUTPUTS = [
    ("irr", "IRR", re.compile(r"^\s*(?:gross\s+|net\s+|levered\s+|equity\s+|sponsor\s+)?IRR\b", re.I), "pct"),
    ("moic", "MOIC", re.compile(r"^\s*(?:gross\s+|net\s+)?(?:MOIC|MoM|multiple\s+of\s+(?:money|invested\s+capital)|cash[- ]on[- ]cash)\b", re.I), "x"),
    ("entry_ev", "Entry EV", re.compile(r"^\s*(?:entry\s+|implied\s+|purchase\s+)?(?:enterprise\s+value|EV)\b(?!\s*/)", re.I), "money"),
    ("purchase_price", "Purchase price", re.compile(r"^\s*(?:total\s+)?purchase\s+price\b", re.I), "money"),
    ("entry_multiple", "Entry multiple", re.compile(r"^\s*(?:entry|purchase|acquisition)\s+(?:EV\s*/\s*EBITDA\s+)?multiple\b|^\s*entry\s+EV\s*/\s*EBITDA", re.I), "x"),
    ("exit_multiple", "Exit multiple", re.compile(r"^\s*exit\s+(?:EV\s*/\s*EBITDA\s+)?multiple\b|^\s*exit\s+EV\s*/\s*EBITDA", re.I), "x"),
    ("equity", "Sponsor equity", re.compile(r"^\s*(?:sponsor|total|cash)?\s*equity\s+(?:contribution|investment|cheque|check)\b", re.I), "money"),
    ("leverage", "Leverage", re.compile(r"^\s*(?:total\s+|senior\s+|net\s+)?(?:debt|leverage)\s*/\s*EBITDA|^\s*(?:total\s+)?leverage\b", re.I), "x"),
    ("exit_ev", "Exit EV", re.compile(r"^\s*exit\s+(?:enterprise\s+value|EV)\b(?!\s*/)", re.I), "money"),
    ("hold", "Hold period", re.compile(r"^\s*(?:hold(?:ing)?\s+period|exit\s+year)\b", re.I), "num"),
]


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _fmt(v, kind):
    if kind == "pct":
        return f"{v * 100:.1f}%" if abs(v) <= 5 else f"{v:.1f}%"
    if kind == "x":
        return f"{v:.2f}x"
    if kind == "money":
        return f"{v:,.1f}"
    return f"{v:,.4g}"


def parse_model(path, filename):
    ext = os.path.splitext(filename.lower())[1]
    if ext not in (".xlsx", ".xlsm"):
        return {"supported": False, "note": "Preview available for .xlsx / .xlsm files only."}
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # corrupt / password protected
        return {"supported": False, "note": f"Could not open workbook: {exc}"}

    sheets, outputs, formulas_missing = [], {}, False
    for ws in wb.worksheets:
        n_rows = n_cols = 0
        for r_idx, row in enumerate(ws.iter_rows(max_row=400, max_col=60, values_only=True), start=1):
            if any(v is not None for v in row):
                n_rows = r_idx
            for c_idx, v in enumerate(row):
                if v is not None:
                    n_cols = max(n_cols, c_idx + 1)
                if not isinstance(v, str) or len(v) > 60:
                    continue
                for key, label, rx, kind in KEY_OUTPUTS:
                    if key in outputs or not rx.search(v):
                        continue
                    # first numeric value to the right on the same row
                    nums = [x for x in row[c_idx + 1: c_idx + 15] if _is_num(x)]
                    if not nums:
                        if any(x is None for x in row[c_idx + 1: c_idx + 4]):
                            formulas_missing = True
                        continue
                    val = nums[0]
                    if kind == "x" and not (0 < val < 100):
                        continue
                    outputs[key] = {"label": label, "cell_label": v.strip(), "value": val,
                                    "display": _fmt(val, kind), "sheet": ws.title, "row": r_idx}
                    break
        sheets.append({"name": ws.title, "rows": n_rows, "cols": n_cols})
    wb.close()
    note = None
    if not outputs and formulas_missing:
        note = "No cached values found - open and save the model in Excel so calculated outputs are stored."
    order = [k for k, *_ in KEY_OUTPUTS]
    return {
        "supported": True,
        "sheets": sheets,
        "outputs": [dict(key=k, **outputs[k]) for k in order if k in outputs],
        "note": note,
    }
